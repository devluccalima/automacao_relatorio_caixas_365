"""Gera a planilha de alterações da carga inicial a partir de carga_inicial/padronizacao_final.xlsx.

Para cada conta do relatório (mesmas regras de INCLUIR_* do .env), lê o cargo, o setor e o
empreendimento atuais no Entra, aplica o de-para e gera:

  carga_inicial/Alteracoes_carga_inicial.xlsx
    - Pendentes:   no formato do Alteracoes.xlsx (só os campos que mudam)
    - Conferência: conta por conta, valor atual e valor novo, para você revisar
    - Sem destino: valores que não estão no de-para nem na lista oficial (não entram)

Uso:
  python ferramentas/carga_inicial.py              gera o arquivo para conferência
  python ferramentas/carga_inicial.py --publicar   gera e envia como Alteracoes.xlsx ao SharePoint
                                                   (recusa se lá já houver linhas pendentes)
"""
import argparse
import sys
import tempfile
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
load_dotenv(RAIZ / ".env", override=True)

from core import alteracoes, analise, leitura_tenant, referencias, sharepoint  # noqa: E402
from core.referencias import chave  # noqa: E402

ENTRADA = RAIZ / "carga_inicial" / "padronizacao_final.xlsx"
SAIDA = RAIZ / "carga_inicial" / "Alteracoes_carga_inicial.xlsx"
ABAS = {"Cargo": "Cargos fora do padrão", "Setor": "Setores", "Empreendimento": "Empreendimentos"}


def ler_de_para() -> dict[str, dict[str, str]]:
    """{campo: {chave(valor atual): valor novo}} a partir das colunas A (atual) e C (padrão)."""
    if not ENTRADA.exists():
        sys.exit(f"Não encontrei {ENTRADA}")
    de_para = {}
    for campo, aba in ABAS.items():
        df = pd.read_excel(ENTRADA, sheet_name=aba, dtype=str).fillna("")
        de_para[campo] = {chave(a): " ".join(n.split()) for a, n in zip(df.iloc[:, 0], df.iloc[:, 2]) if a and n}
    return de_para


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--publicar", action="store_true", help="envia como Alteracoes.xlsx ao SharePoint")
    args = ap.parse_args()

    refs = referencias.carregar(RAIZ / "config")
    oficiais = {"Cargo": refs.nomes_cargos, "Setor": refs.setores, "Empreendimento": refs.empreendimentos}
    de_para = ler_de_para()

    print("Lendo contas do Entra ID...")
    skus = leitura_tenant.ler_skus()
    contas = leitura_tenant.ler_contas(skus)
    caixas = analise.analisar(contas, refs).caixas  # mesmas regras de inclusão do relatório

    no_relatorio = {u.lower() for u in caixas["User principal name"]}
    campos_conta = {"Cargo": "cargo", "Setor": "setor", "Empreendimento": "empreendimento"}

    pendentes, conferencia, sem_destino = [], [], []
    for c in sorted(contas, key=lambda x: x["nome"].casefold()):
        if c["upn"].lower() not in no_relatorio:
            continue
        linha = {col: "" for col in alteracoes.COLUNAS}
        linha["E-mail"] = c["upn"]
        conf = {"E-mail": c["upn"], "Display name": c["nome"]}
        muda = False
        for campo, chave_conta in campos_conta.items():
            atual = c[chave_conta]  # valor bruto do Entra, com eventuais espaços sobrando
            novo = ""
            if atual.strip():
                if chave(atual) in de_para[campo]:
                    novo = de_para[campo][chave(atual)]
                elif chave(atual) in oficiais[campo]:
                    novo = oficiais[campo][chave(atual)]
                else:
                    sem_destino.append({"E-mail": c["upn"], "Campo": campo, "Valor atual": atual})
            conf[f"{campo} atual"] = atual
            conf[f"{campo} novo"] = novo if novo and novo != atual else ""
            if novo and novo != atual and novo == " ".join(atual.split()):
                conf.setdefault("Nota", "")
                conf["Nota"] += f"{campo}: só remove espaço sobrando. "
            if novo and novo != atual:
                linha[campo] = novo
                muda = True
        if muda:
            linha["Observação"] = "Carga inicial"
            pendentes.append(linha)
            conferencia.append(conf)

    SAIDA.parent.mkdir(exist_ok=True)
    alteracoes.montar(pendentes, [], refs, SAIDA)
    wb = load_workbook(SAIDA)
    for titulo, dados in (("Conferência", conferencia), ("Sem destino", sem_destino)):
        ws = wb.create_sheet(titulo)
        df = pd.DataFrame(dados)
        if df.empty:
            ws.append(["Nada aqui"])
            continue
        ws.append(list(df.columns))
        for linha in df.itertuples(index=False):
            ws.append(list(linha))
        for cel in ws[1]:
            cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="1F4E78")
        for i in range(len(df.columns)):
            ws.column_dimensions[chr(65 + i)].width = 34
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(SAIDA)

    contagem = {c: sum(1 for p in pendentes if p[c]) for c in ("Cargo", "Setor", "Empreendimento")}
    print(f"\n{len(caixas)} contas no relatório; {len(pendentes)} com alguma mudança")
    print(f"  cargos: {contagem['Cargo']}  setores: {contagem['Setor']}  empreendimentos: {contagem['Empreendimento']}")
    print(f"  valores sem destino (não entram): {len(sem_destino)}")
    print(f"Arquivo para conferência: {SAIDA}")

    if args.publicar:
        if not sharepoint.configurado():
            sys.exit("SharePoint não configurado: copie o arquivo para alteracoes/Alteracoes.xlsx manualmente.")
        with tempfile.TemporaryDirectory() as t:
            existente = sharepoint.baixar_arquivo(alteracoes.NOME_ARQUIVO, Path(t) / "a.xlsx")
            historico = []
            if existente:
                pend, historico = alteracoes.ler(existente)
                if pend:
                    sys.exit(f"O {alteracoes.NOME_ARQUIVO} do SharePoint já tem {len(pend)} linhas pendentes. "
                             "Resolva ou apague essas linhas antes de publicar a carga inicial.")
            destino = alteracoes.montar(pendentes, historico, refs, Path(t) / "pub.xlsx")
            envio = sharepoint.enviar_arquivo(destino, alteracoes.NOME_ARQUIVO)
        print(f"Publicado no SharePoint: {envio['url']}")


if __name__ == "__main__":
    main()