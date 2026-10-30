"""Push: aplica no Exchange as linhas da planilha Alteracoes.xlsx.

A planilha fica na mesma pasta do SharePoint do relatório (ou em alteracoes/ se o SharePoint
não estiver configurado) e tem quatro abas:

- Pendentes: uma linha por conta. E-mail obrigatório; Cargo, Setor, Empreendimento, Nome próprio,
  Apelido e Nome a apresentar são opcionais (vazio = não mexe). "Observação" é livre para quem
  preenche; "Status" é escrito pelo script.
- Sugestões: valores fora do padrão para os quais existe um oficial parecido. Escreva S em
  "Aprovar (S)" (e, se quiser, troque a sugestão) e a linha vira alteração na próxima execução.
- Histórico: o que já foi aplicado, com data e resultado.
- Listas (oculta): alimenta as listas suspensas de Cargo, Setor e Empreendimento.

Linhas aplicadas (OK ou SEM_MUDANCA) saem de Pendentes e vão para o Histórico. Linhas com erro,
inválidas, simuladas ou além do limite da execução ficam em Pendentes com o Status explicando.
"""
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core import exchange, sharepoint
from core.referencias import Referencias, chave

RAIZ = Path(__file__).resolve().parents[1]
NOME_ARQUIVO = "Alteracoes.xlsx"
PASTA_LOCAL = RAIZ / "alteracoes"

CAMPOS = ["Cargo", "Setor", "Empreendimento", "Nome próprio", "Apelido", "Nome a apresentar"]
COLUNAS = ["E-mail", *CAMPOS, "Observação", "Status"]
COLUNAS_HIST = ["Data", "Modo", "E-mail", "Resultado", "Detalhe", "Observação"]
COLUNAS_SUG = ["Aprovar (S)", "E-mail", "Display name", "Campo", "Valor atual", "Sugestão", "Semelhança"]
OBS_AUTO = "Autocorreção"
OBS_SUG = "Sugestão aprovada"
PARA_PS = {"Cargo": "Cargo", "Setor": "Setor", "Empreendimento": "Empreendimento",
           "Nome próprio": "NomeProprio", "Apelido": "Apelido", "Nome a apresentar": "NomeApresentar"}
SINONIMOS = {"email": "E-mail", "e-mail": "E-mail", "user principal name": "E-mail",
             "departamento": "Setor", "escritório": "Empreendimento", "escritorio": "Empreendimento",
             "nome proprio": "Nome próprio", "nome a apresentar": "Nome a apresentar",
             "display name": "Nome a apresentar"}
LINHAS_COM_LISTA = 3000

CABECALHO = PatternFill("solid", fgColor="1F4E78")
FONTE = Font(name="Arial", size=10)


# ------------------------------------------------------------------ planilha

def _texto(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def ler(caminho: Path) -> tuple[list[dict], list[list]]:
    wb = load_workbook(caminho, data_only=True)
    ws = wb["Pendentes"] if "Pendentes" in wb.sheetnames else wb.worksheets[0]
    cab = [_texto(c.value) for c in ws[1]]
    nomes = [SINONIMOS.get(c.casefold(), next((x for x in COLUNAS if x.casefold() == c.casefold()), c))
             for c in cab]
    pendentes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        linha = {n: _texto(v) for n, v in zip(nomes, row) if n in COLUNAS}
        if any(linha.get(c) for c in ["E-mail", *CAMPOS]):
            pendentes.append({c: linha.get(c, "") for c in COLUNAS})
    historico = []
    if "Histórico" in wb.sheetnames:
        historico = [list(r) for r in wb["Histórico"].iter_rows(min_row=2, values_only=True) if any(r)]
    return pendentes, historico


def ler_sugestoes(caminho: Path) -> list[dict]:
    wb = load_workbook(caminho, data_only=True)
    if "Sugestões" not in wb.sheetnames:
        return []
    ws = wb["Sugestões"]
    cab = [_texto(c.value) for c in ws[1]]
    sugestoes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        linha = {n: _texto(v) for n, v in zip(cab, row) if n in COLUNAS_SUG}
        if linha.get("E-mail"):
            sugestoes.append({c: linha.get(c, "") for c in COLUNAS_SUG})
    return sugestoes


def montar(pendentes: list[dict], historico: list[list], refs: Referencias, destino: Path,
           sugestoes: list[dict] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pendentes"
    ws.append(COLUNAS)
    for p in pendentes:
        ws.append([p.get(c, "") for c in COLUNAS])
    larguras = [42, 38, 30, 30, 18, 22, 34, 34, 70]
    for i, w in enumerate(larguras):
        ws.column_dimensions[chr(65 + i)].width = w

    wsug = wb.create_sheet("Sugestões")
    wsug.append(COLUNAS_SUG)
    for sg in sugestoes or []:
        wsug.append([sg.get(c, "") for c in COLUNAS_SUG])
    for i, w in enumerate([12, 42, 34, 16, 40, 40, 12]):
        wsug.column_dimensions[chr(65 + i)].width = w
    dv_sug = DataValidation(type="list", formula1='"S"', allow_blank=True)
    dv_sug.add(f"A2:A{LINHAS_COM_LISTA}")
    wsug.add_data_validation(dv_sug)

    wh = wb.create_sheet("Histórico")
    wh.append(COLUNAS_HIST)
    for h in historico:
        wh.append(list(h)[:len(COLUNAS_HIST)])
    for i, w in enumerate([17, 11, 42, 14, 90, 34]):
        wh.column_dimensions[chr(65 + i)].width = w

    wl = wb.create_sheet("Listas")
    listas = {"A": sorted(refs.nomes_cargos.values(), key=str.casefold),
              "B": sorted(refs.setores.values(), key=str.casefold),
              "C": sorted(refs.empreendimentos.values(), key=str.casefold)}
    for col, (titulo, valores) in zip("ABC", [("Cargo", listas["A"]), ("Setor", listas["B"]),
                                                 ("Empreendimento", listas["C"])]):
        wl[f"{col}1"] = titulo
        for i, v in enumerate(valores, start=2):
            wl[f"{col}{i}"] = v
    wl.sheet_state = "hidden"
    for col_lista, col_pend in (("A", "B"), ("B", "C"), ("C", "D")):
        n = len(listas[col_lista])
        if not n:
            continue
        dv = DataValidation(type="list", formula1=f"=Listas!${col_lista}$2:${col_lista}${n + 1}",
                            allow_blank=True, showErrorMessage=True, errorStyle="stop",
                            errorTitle="Valor fora da lista",
                            error="Escolha um valor da lista (ou deixe em branco para não alterar).")
        dv.add(f"{col_pend}2:{col_pend}{LINHAS_COM_LISTA}")
        ws.add_data_validation(dv)

    for aba in (ws, wsug, wh):
        for row in aba.iter_rows():
            for cel in row:
                cel.font = FONTE
        for cel in aba[1]:
            cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cel.fill = CABECALHO
        aba.freeze_panes = "A2"
    wb.save(destino)
    return destino


# ------------------------------------------------------------------ validação

def validar(linha: dict, refs: Referencias) -> tuple[dict | None, str]:
    """Devolve (linha pronta para o PowerShell, "") ou (None, motivo)."""
    email = linha["E-mail"].lower()
    if "@" not in email:
        return None, "E-mail vazio ou inválido"
    if not any(linha[c] for c in CAMPOS):
        return None, "Nenhum campo preenchido para alterar"
    saida = {"Email": email}
    for campo, oficiais in (("Cargo", refs.nomes_cargos), ("Setor", refs.setores),
                            ("Empreendimento", refs.empreendimentos)):
        valor = linha[campo]
        if not valor:
            continue
        if oficiais:
            if chave(valor) not in oficiais:
                return None, f"{campo} '{valor}' não está na lista oficial"
            valor = oficiais[chave(valor)]  # grafia oficial
        saida[PARA_PS[campo]] = valor
    for campo in ("Nome próprio", "Apelido", "Nome a apresentar"):
        if linha[campo]:
            saida[PARA_PS[campo]] = linha[campo]
    return saida, ""


# ------------------------------------------------------------------ ciclo

def _chave_linha(p: dict) -> tuple:
    return tuple(chave(p.get(c, "")) for c in ["E-mail", *CAMPOS])


def _obter(tmp: Path) -> Path | None:
    local = tmp / NOME_ARQUIVO
    if sharepoint.configurado():
        return sharepoint.baixar_arquivo(NOME_ARQUIVO, local)
    arq = PASTA_LOCAL / NOME_ARQUIVO
    if arq.exists():
        local.write_bytes(arq.read_bytes())
        return local
    return None


def _gravar(arquivo: Path) -> None:
    if sharepoint.configurado():
        envio = sharepoint.enviar_arquivo(arquivo, NOME_ARQUIVO)
        if envio["copia"]:
            print(f"  ATENÇÃO: {NOME_ARQUIVO} estava aberto; o resultado foi gravado em {envio['nome']}. "
                  "As linhas já aplicadas podem aparecer de novo e serão marcadas SEM_MUDANCA.")
    else:
        PASTA_LOCAL.mkdir(exist_ok=True)
        (PASTA_LOCAL / NOME_ARQUIVO).write_bytes(arquivo.read_bytes())


def _incorporar(pendentes: list[dict], linhas: list[dict], observacao: str, mesclar: bool = False) -> int:
    """linhas: {"E-mail", "Campo", "Valor"}. Junta por e-mail numa linha nova.

    Se o e-mail já tinha linha pendente: com mesclar=True, preenche o campo nela quando estiver vazio
    (usado nas sugestões aprovadas); com mesclar=False, não mexe (autocorreção não passa por cima de
    uma alteração feita por alguém)."""
    existentes = {p["E-mail"].lower(): p for p in pendentes}
    novas: dict[str, dict] = {}
    for l in linhas:
        email = l["E-mail"].lower()
        if email in existentes:
            if mesclar and not existentes[email].get(l["Campo"]):
                existentes[email][l["Campo"]] = l["Valor"]
                obs = existentes[email].get("Observação", "")
                if observacao not in obs:
                    existentes[email]["Observação"] = f"{obs} + {observacao}" if obs else observacao
                novas.setdefault(email, existentes[email])
            continue
        linha = novas.setdefault(email, {**{c: "" for c in COLUNAS}, "E-mail": l["E-mail"], "Observação": observacao})
        linha[l["Campo"]] = l["Valor"]
    pendentes.extend(n for e, n in novas.items() if e not in existentes)
    return len(novas)


def processar(refs: Referencias, simular: bool, limite: int | None = None,
              autocorrecoes: list[dict] | None = None, sugestoes: list[dict] | None = None) -> dict:
    """Lê a planilha, incorpora sugestões aprovadas e autocorreções, valida, aplica (ou simula)
    e grava o resultado. Devolve um resumo.

    autocorrecoes: [{"E-mail", "Campo", "Valor"}] viram linhas em Pendentes (Observação "Autocorreção").
    sugestoes: lista nova para a aba Sugestões; None mantém a que está na planilha.
    """
    limite = limite if limite is not None else int(os.getenv("LIMITE_ALTERACOES", "50"))
    agora = datetime.now()
    resumo = {"pendentes": 0, "aplicadas": 0, "sem mudança": 0, "simuladas": 0, "erros": 0,
              "inválidas": 0, "aguardando": 0, "autocorreções": 0, "sugestões aprovadas": 0,
              "sugestões": 0, "detalhes": []}

    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        arquivo = _obter(tmp)
        if arquivo is None:
            print(f"  {NOME_ARQUIVO} não existe; criando a planilha.")
            pendentes, historico, sug_atuais = [], [], []
        else:
            pendentes, historico = ler(arquivo)
            sug_atuais = ler_sugestoes(arquivo)
        originais = {_chave_linha(p) for p in pendentes}

        aprovadas = [sg for sg in sug_atuais if sg["Aprovar (S)"].strip().upper() in ("S", "SIM")]
        resumo["sugestões aprovadas"] = _incorporar(
            pendentes, [{"E-mail": sg["E-mail"], "Campo": sg["Campo"], "Valor": sg["Sugestão"]} for sg in aprovadas
                        if sg["Campo"] in CAMPOS], OBS_SUG, mesclar=True)
        resumo["autocorreções"] = _incorporar(pendentes, autocorrecoes or [], OBS_AUTO)

        # Sugestão some só se aquele campo da conta já tiver um valor pendente
        pend_por_email = {p["E-mail"].lower(): p for p in pendentes}
        base_sug = sugestoes if sugestoes is not None else [sg for sg in sug_atuais if sg not in aprovadas]
        nova_sug = [{**sg, "Aprovar (S)": ""} for sg in base_sug
                    if not pend_por_email.get(sg["E-mail"].lower(), {}).get(sg["Campo"])]
        resumo["sugestões"] = len(nova_sug)
        resumo["pendentes"] = len(pendentes)

        if not pendentes:
            # Regrava para manter as listas suspensas e as sugestões atualizadas
            _gravar(montar([], historico, refs, tmp / "novo.xlsx", nova_sug))
            return resumo

        validas, vistos = [], set()
        for p in pendentes:
            pronta, motivo = validar(p, refs)
            if pronta and pronta["Email"] in vistos:
                pronta, motivo = None, "E-mail repetido em outra linha; junte as duas"
            if pronta is None:
                p["Status"] = f"Inválida ({agora:%d/%m %H:%M}): {motivo}"
                resumo["inválidas"] += 1
                continue
            vistos.add(pronta["Email"])
            if len(validas) >= limite:
                p["Status"] = f"Aguardando: limite de {limite} contas por execução"
                resumo["aguardando"] += 1
                continue
            validas.append((p, pronta))

        resultado = {}
        if validas:
            modo = "simulação" if simular else "gravando"
            print(f"  {len(validas)} contas para o Exchange ({modo})...")
            df = exchange.aplicar([pr for _, pr in validas], simular)
            resultado = {r.Email.lower(): (r.Resultado, r.Detalhe) for r in df.itertuples()}

        restantes = []
        for p in pendentes:
            par = next((pr for lp, pr in validas if lp is p), None)
            if par is None:
                restantes.append(p)
                continue
            res, det = resultado.get(par["Email"], ("ERRO", "Sem resposta do script do Exchange"))
            resumo["detalhes"].append((par["Email"], res, det))
            registro = [f"{agora:%d/%m/%Y %H:%M}", "Real", par["Email"], res, det, p["Observação"]]
            if res == "SIMULADO":
                resumo["simuladas"] += 1
                p["Status"] = f"Simulado ({agora:%d/%m %H:%M}): {det}"
                restantes.append(p)
            elif res == "ERRO":
                resumo["erros"] += 1
                p["Status"] = f"ERRO ({agora:%d/%m %H:%M}): {det}"
                restantes.append(p)
                if not simular:
                    historico.append(registro)
            else:
                resumo["aplicadas" if res == "OK" else "sem mudança"] += 1
                historico.append(registro)

        # Junta linhas que alguém acrescentou enquanto o Exchange processava
        pasta_atual = tmp / "atual"
        pasta_atual.mkdir()
        atual = _obter(pasta_atual)
        if atual:
            novos, _ = ler(atual)
            restantes += [n for n in novos if _chave_linha(n) not in originais]

        _gravar(montar(restantes, historico, refs, tmp / "novo.xlsx", nova_sug))
    return resumo